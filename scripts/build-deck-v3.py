# v2 덱 → v3: 앱 슬라이드 갱신 + 근거 프리셋·설계 규칙·계약자 조정·현행 제도 점검 ①② 추가
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC = "C:/00 App Project/flexible_whole_life/docs/tkLeen_설계형종신보험_제안슬라이드_v2.pptx"
DST = "C:/00 App Project/flexible_whole_life/docs/tkLeen_설계형종신보험_제안슬라이드_v3.pptx"
SHOTS = "docs/slides/"   # 앱 스크린샷(headless Chromium, 2x). 저장소 루트에서 실행
NAVY, SKY, GRAY, INK, LINE, PANEL, RUST = "1B2845", "4A90C2", "5A6673", "0A0A0A", "E3E6EA", "FBFCFD", "C2704A"

prs = Presentation(SRC)
layout = prs.slides[0].slide_layout
logo_blob = next(sh.image.blob for sh in prs.slides[0].shapes if sh.shape_type == 13)
import io

def rgb(h): return RGBColor.from_string(h)

def text(slide, x, y, w, h, s, size=12, bold=False, color=INK, font="Pretendard", align=None, anchor=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0); tf.margin_top = tf.margin_bottom = Emu(0)
    if anchor: tf.vertical_anchor = anchor
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align: p.alignment = align
        r = p.add_run(); r.text = line
        r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color)
    return tb

def rect(slide, x, y, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    if line: sh.line.color.rgb = rgb(line); sh.line.width = Pt(0.75)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh

def hline(slide, x, y, w, color=NAVY):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = rgb(color); ln.line.width = Pt(0.75)
    return ln

def chrome(slide, label, title, subtitle=None):
    text(slide, 0.70, 0.42, 6.0, 0.30, label, 10, False, SKY, "JetBrains Mono")
    text(slide, 0.70, 0.72, 11.93, 0.60, title, 26, True, NAVY)
    if subtitle: text(slide, 0.70, 1.32, 11.93, 0.35, subtitle, 13, False, GRAY)
    hline(slide, 0.70, 1.72, 11.93)
    slide.shapes.add_picture(io.BytesIO(logo_blob), Inches(0.70), Inches(6.88), height=Inches(0.34))
    text(slide, 1.04, 6.90, 1.5, 0.30, "tkLeen", 11, True, INK, "Fraunces")
    slide.shapes.add_textbox(Inches(8.63), Inches(6.90), Inches(4.0), Inches(0.30)).name = "FOOTER_PAGE"

def table(slide, x, y, w, cols, rows, head_size=12, body_size=11.0, row_h=0.32):
    n_rows, n_cols = len(rows) + 1, len(cols)
    gs = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(row_h * n_rows))
    tb = gs.table
    tblPr = tb._tbl.tblPr
    for child in list(tblPr): tblPr.remove(child)          # 기본 스타일 제거
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    total = sum(c[1] for c in cols)
    for i, (name, wf) in enumerate(cols): tb.columns[i].width = Inches(w * wf / total)
    def fill_cell(c, s, size, bold, color, bg=None):
        c.margin_left = c.margin_right = Emu(76200); c.margin_top = c.margin_bottom = Emu(50800)
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; r = p.add_run(); r.text = s
        r.font.name = "Pretendard"; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color)
        if bg: c.fill.solid(); c.fill.fore_color.rgb = rgb(bg)
        else: c.fill.background()
        # 얇은 회색 테두리
        tcPr = c._tc.get_or_add_tcPr()
        for side in ("lnL", "lnR", "lnT", "lnB"):
            ln = etree.SubElement(tcPr, qn("a:" + side), w="6350"); sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=LINE)
    tb.rows[0].height = Inches(0.36)
    for ri in range(1, n_rows): tb.rows[ri].height = Inches(row_h)
    for i, (name, _) in enumerate(cols): fill_cell(tb.cell(0, i), name, head_size, True, NAVY, "FAFAF7")
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row): fill_cell(tb.cell(ri, ci), val, body_size, ci == 0, NAVY if ci == 0 else INK)
    return gs

def note(slide, y, label, body, accent=NAVY, h=1.0):
    rect(slide, 0.70, y, 11.93, h, PANEL, LINE); rect(slide, 0.70, y, 0.06, h, accent)
    text(slide, 0.95, y + 0.12, 4.0, 0.25, label, 9, False, NAVY, "JetBrains Mono")
    text(slide, 0.95, y + 0.40, 11.43, h - 0.5, body, 11.5, False, INK)

# ---------------- 기존 슬라이드 갱신 ----------------
def find_text(slide, startswith):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith(startswith): return sh
    return None
def set_text(sh, s):
    p = sh.text_frame.paragraphs[0]
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    p.runs[0].text = s
    for extra in sh.text_frame.paragraphs[1:]: extra._p.getparent().remove(extra._p)

# slide 8: 앱 소개 — 부제·이미지 교체
s8 = prs.slides[7]
set_text(find_text(s8, "입력 → 경로 편집"), "조건 입력 → 근거 있는 프리셋 → 계약자가 그래프에서 조정 → 규칙 검증 → 보험료·환급금 산출 → 대안 비교")
for sh in list(s8.shapes):
    if sh.shape_type == 13 and sh.width > Inches(2): sh._element.getparent().remove(sh._element)
s8.shapes.add_picture(SHOTS + "design.png", Inches(1.55), Inches(1.9), height=Inches(4.75))

# slide 20: 개발 단계 — 현재 상태 반영
s20 = prs.slides[19]
set_text(find_text(s20, "Next.js 14"), "Next.js 15 + 순수 TypeScript 엔진 + 브라우저 보관함(JSON·공유 링크) · 2026-09 현재 1~3단계 완료, 자동 테스트 139개·골든값 재현")
set_text(find_text(s20, "프리셋 + 계산 엔진"), "완료 · 프리셋 6종 + 이중탈퇴 계산 엔진 + Excel 검산 시트(SUMPRODUCT 수식) + 수식 설명(?) 팝업")
set_text(find_text(s20, "추천 엔진, 시나리오"), "완료 · 니즈·HLV 추천, 같은 예산 3안 비교, 제안서 인쇄, 공유 링크, 설정(가정 세트·envelope)")
set_text(find_text(s20, "재설계 모드(기존"), "완료 · 재설계 모드(원계약→예산→재설계 캔버스→R01~R03·전환 손실 비교) + 근거 기반 프리셋·설계 규칙 팝업")
for sh in s20.shapes:
    if sh.shape_type == 1 and abs(sh.top / 914400 - 6.05) < 0.01: sh.height = Inches(0.75)
set_text(find_text(s20, "회사별 기초율"), "다음 · 회사별 기초율 업로드, 인가 문서(산출방법서·약관 초안) 자동 생성, 사업비·준비금 민감도, 계약관리 API 규격, 로그인·서버 저장")

# ---------------- 새 슬라이드 ----------------
new_slides = []
def new(label, title, subtitle=None):
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    chrome(s, label, title, subtitle); new_slides.append(s); return s

# A. 근거 기반 프리셋
sA = new("PART 1 · APP ②", "근거 있는 프리셋: 조건을 넣으면 필요액 곡선이 보장이 된다", "프리셋은 회사가 정한 모양이 아니라 계약자의 조건에서 수식으로 나온 출발점이다. 공통 조건은 한 번만 입력한다")
table(sA, 0.70, 1.95, 11.93, [("프리셋", 1.3), ("필요액 산정(이론·수식)", 5.4), ("계약자 조건", 3.2), ("기준보험금 제안", 2.0)], [
    ["자녀연령형", "유족 생활비 현가 Y×70%×a(막내 독립까지) + 1억×독립 전 자녀 수 + 부채 잔액 + 정리자금 − 유동자산·기존 보장. 독립 뒤엔 정리자금만", "자녀 나이·연소득·부채(공통)", "현재 필요액"],
    ["부채상환형", "대출 잔액 B_t (원리금균등 닫힌식·원금균등·만기일시) + 정리자금", "잔액·만기(공통), 금리, 상환방식", "잔액 + 정리자금"],
    ["은퇴증액형", "은퇴 후 = 배우자 월 생활비×12×a(기대여명 e_s, 경험생명표) + 정리자금 − 은퇴 자산. 배수 R = 은퇴 후 ÷ 은퇴 전(니즈)", "은퇴시기(공통)·배우자 나이·월 생활비·은퇴 자산", "니즈"],
    ["단체보험보완형", "재직 중 g = 1 − 단체보험 ÷ 니즈, 퇴직(은퇴시기)과 함께 100%", "단체보험 보험금·은퇴시기(공통)", "니즈"],
    ["상속준비형", "순자산 A_0(1+g)^t 의 상속세(일괄공제 5억·배우자공제·10~50% 누진) ÷ 70세 상속세", "순자산·증가율·배우자·자녀 수(공통)", "70세 상속세"],
    ["평준형", "니즈(생활비·교육비·부채·정리자금 − 자산·기존 보장) 또는 HLV(소득 현가)", "연소득·유동자산·기존 보장(공통)", "니즈 / HLV"],
], body_size=10.5, row_h=0.5)
note(sA, 5.55, "곡선 → 보장 스케줄", "필요액 곡선을 그대로 쓰지 않고 설계 규칙에 맞춘다: 증액은 필요 시점에 맞춰 미리 시작하고(매년 최대 10%), 감액은 매년 10%씩 늦게 내려 보장이 필요액 아래로 가지 않는다. 초기 5년 고정·70세 이후 증액 없음·3배 상한·20% 하한을 같은 절차에서 적용한다. 카드의 \"?\"가 근거와 계산 수치를, \"입력\"이 조건 팝업을 연다.", h=1.15)

# B. 설계 규칙
sB = new("PART 1 · APP ③", "설계 규칙: 계약자가 그려도 지켜지는 원칙", "규칙은 코드에 숨기지 않고 화면의 \"설계 규칙\" 팝업으로 보여준다 · 프리셋·그래프 편집·재설계가 모두 같은 규칙을 쓴다")
table(sB, 0.70, 1.95, 7.3, [("코드", 0.9), ("규칙", 1.6), ("내용(기본값, 설정에서 변경)", 4.8)], [
    ["E01", "초기 고정", "가입 후 5년은 보험금 변경 불가(역선택 대기기간)"],
    ["1칸", "단위·범위", "1칸 = 기준보험금의 10%. 연령 A에서 ±(A − 직전 변경점)칸, 매년 최대 1칸"],
    ["E02", "연 증가율", "20% 이내(단, 1칸 10%는 항상 허용)"],
    ["E03", "증액 종료", "70세 이후 증액 불가(감액은 가능)"],
    ["E04·E05", "상·하한", "최대 3배, 기준보험금의 20% 이상"],
    ["E06·E07", "금액", "보험금 1천만원 이상, 기준보험금 10억 이하(심사 한도)"],
    ["E08", "축하금", "해당 연령 보험금의 10%, 누계 ≤ 그때까지 낸 보험료"],
    ["곡선 적합", "프리셋", "증액 앞당김·감액 지연·5년 고정·70세 후 증액 없음·상하한"],
], body_size=10.5, row_h=0.42)
sB.shapes.add_picture(SHOTS + "rules.png", Inches(8.25), Inches(1.95), width=Inches(4.38))
note(sB, 5.85, "왜 규칙을 보여 주는가", "설계 자유도가 높을수록 감독기관은 \"복잡한 상품\"으로 읽는다. 허용 범위(envelope)를 산출방법서와 화면에 같은 문장으로 두면 인가 심사·설명의무·시스템 검증이 하나의 규격을 공유한다.", h=0.85)

# C. 계약자 조정
sC = new("PART 1 · APP ④", "계약자가 직접 조정한다: 프리셋 위에 자기 상황을 덧그린다", "회사가 정한 구조를 고르는 것이 아니라, 근거 있는 출발점을 자기 손으로 고친다")
items = [
    ("드래그", "연령을 누르고 위아래로 끌면 그 연령부터 같은 폭만큼 이동하고, 앞쪽에는 매년 한 칸씩 잇는 램프가 놓인다. 프리셋 모양은 유지된다"),
    ("범위", "직전 변경점이 정한 수준에서 지난 연수만큼 ±칸. 올리든 내리든 매년 최대 10%"),
    ("더블클릭", "그 연령 이후를 그 값으로 평탄화. 직선 위에 남는 변경점(마름모)은 자동으로 지운다"),
    ("되돌리기·초기화", "한 번의 드래그가 한 단계. 초기화는 1억·표준 평준형으로"),
    ("예산", "월 보험료 슬라이더는 모양은 두고 기준보험금(1천만원 단위)만 바꾼다"),
    ("검증", "E01~E08 위반은 배지로 즉시 표시, 자동 수정 제안"),
]
y = 1.95
for i, (k, v) in enumerate(items):
    rect(sC, 0.70, y, 0.5, 0.5, SKY); text(sC, 0.70, y + 0.1, 0.5, 0.3, str(i + 1), 13, True, "FFFFFF", "JetBrains Mono", PP_ALIGN.CENTER)
    text(sC, 1.35, y - 0.02, 1.7, 0.3, k, 13, True, NAVY); text(sC, 1.35, y + 0.26, 5.4, 0.45, v, 10.5, False, INK)
    y += 0.78
sC.shapes.add_picture(SHOTS + "evidence.png", Inches(7.25), Inches(1.95), width=Inches(5.38))
text(sC, 7.25, 5.4, 5.38, 0.6, "근거 팝업: 이론·표준 모양·조건·현재 입력으로 계산한 필요액과 수식 링크. 설계사가 설명의무를 이행하는 화면이자 계약자가 스스로 확인하는 화면", 10, False, GRAY)

# D1. 현행 제도 점검 ① 상품·판매
sD1 = new("04 · REGULATION ②", "현행 제도 점검 ①: 상품 인가·판매 — 2026년 기준", "설계 자유도가 높은 상품이 현행 제도에서 부딪히는 지점과 대응")
table(sD1, 0.70, 1.95, 11.93, [("제도", 1.7), ("현행 내용", 3.6), ("설계형에서 문제될 부분", 3.3), ("해결 방안", 3.33)], [
    ["기초서류 신고", "보험업법 제127조: 새로운 위험담보·보험료 산출방식 변경 등은 사전 신고, 보험개발원 요율 검증. 감독규정 2026.1.2 개정 시행", "연령별 보험금 벡터가 계약마다 달라 산출방법서를 '하나의 표'로 쓸 수 없다. 심사관은 허용 범위를 요구", "산출방법서를 산식(1년 블록 현가 합·envelope E01~E08)으로 정의하고 프리셋 6종을 대표 사례로 첨부. Excel 검산 시트를 검증 자료로"],
    ["금소법 6대 판매원칙", "적합성·적정성·설명의무·불공정영업 금지·부당권유 금지·광고 규제. 위반 시 청약철회·위법계약해지권", "'계약자가 설계'하면 적합성 판단과 설명의 주체가 흐려진다. 복잡한 상품으로 분류될 위험", "조건 입력(적합성 기록) → 근거 팝업(설명) → 규칙 검증(적정성)을 앱 흐름에 내장하고 제안서에 남긴다. 프리셋을 기본으로 두고 직접 편집은 한 단계 뒤에"],
    ["체증형 종신 소비자경보", "금감원: 체증형은 보험료가 비싸고 설명이 부족, 갈아타기 유도·무·저해지 결합 시 손실 위험 경고", "증액 경로는 곧 체증형. 같은 경보 대상이 될 수 있다", "같은 예산 3안 비교(평준·정기+종신·설계형)를 필수 화면으로, 증액은 필요액 근거가 있을 때만 제안. 저해지는 표준형 뒤에 선택지로"],
    ["판매수수료 개편", "2026.1 수수료 공시·5등급 비교, 2026.7 GA에도 1200% 룰, 첫해 집중 → 7년 분급", "초기 보험금이 낮은 설계(단체보험보완·상속준비)는 첫해 수수료가 낮아 대면채널 유인이 약하다", "수수료 기준액을 '초기 5년 고정 보험금'이 아닌 '스케줄 현가 기준'으로 정의해 신고. 유지 수수료 비중이 큰 분급 체계와 오히려 잘 맞는다"],
], body_size=10, row_h=0.95)

# D2. 현행 제도 점검 ② 계약변경·회계·세제
sD2 = new("04 · REGULATION ③", "현행 제도 점검 ②: 계약 변경·회계·세제·유동화", "약관·회계·세법이 전제하는 '고정 보험금'을 '약정된 스케줄'로 바꿔 읽게 하는 것이 핵심")
table(sD2, 0.70, 1.95, 11.93, [("제도", 1.7), ("현행 내용", 3.6), ("설계형에서 문제될 부분", 3.3), ("해결 방안", 3.33)], [
    ["표준약관 계약 변경", "계약내용 변경은 회사 승낙. 감액은 해지로 보아 환급금 지급, 증액은 계약 전 알릴의무·심사", "매년 바뀌는 보험금을 '변경'으로 보면 해마다 승낙·고지·환급 절차가 생긴다", "청약 시 확정한 스케줄은 변경이 아니라 약정(체증·체감형과 동일). 재설계 옵션을 약관에 내재해 재설계도 옵션 행사로. 확정 경로 밖 증액만 재고지"],
    ["무·저해지 계리가정", "IFRS17 계리가정 가이드라인: 해지율은 로그-선형 원칙모형, 다른 기법은 차이·영향 공시(2025 재무제표부터)", "저해지 옵션 결합 시 스케줄별 해지율 가정과 공시 부담 급증", "표준형 환급 구조로 인가, 저해지는 후속 특약. 스케줄은 해지율 가정과 독립적으로 두고 CSM 조정으로 처리"],
    ["세제", "보험차익 비과세: 월적립식 5년 납·10년 유지·월 150만 한도, 순수보장성은 한도 없음. 계약자=피보험자 사망보험금은 간주상속재산", "해지환급률·축하금(생존급부)이 커지면 저축성으로 분류돼 한도 적용. 상속세 재원 설계는 보험금 자체가 과세되는 점 반영 필요", "축하금 누계 ≤ 납입보험료(E08)로 보장성 유지. 상속준비형은 세액 ÷ (1 − 한계세율)로 기준보험금 상향 안내. 계약자·수익자 구조 안내를 제안서에"],
    ["사망보험금 유동화", "2025.10 시행: 55세 이상·금리확정형·9억 이하 종신의 사망보험금 최대 90%를 연금·서비스로 전환(2026.1 전 생보사)", "후기 보험금이 큰 설계(은퇴증액·상속준비)일수록 유동화 가치가 커 상품 취지와 충돌할 수 있다", "유동화를 설계형의 '은퇴 후 감액 재설계'와 연결: 재설계 화면에서 연금 전환액을 함께 제시. 금리확정형 요건을 상품 규격에 반영"],
], body_size=10, row_h=0.95)

# ---------------- 순서 조정: A·B·C는 슬라이드 8 뒤, D1·D2는 슬라이드 15 뒤 ----------------
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
def move_after(slide, after_index):
    el = next(x for x in sldIdLst if x.get("id") == str(slide.slide_id) or x.rId == prs.part.relate_to(slide.part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"))
new_ids = ids[-5:]           # A B C D1 D2
base = ids[:-5]
order = base[:8] + new_ids[:3] + base[8:15] + new_ids[3:] + base[15:]
for x in list(sldIdLst): sldIdLst.remove(x)
for x in order: sldIdLst.append(x)

# 푸터 페이지 번호 갱신
N = len(prs.slides)
for i, s in enumerate(prs.slides, 1):
    done = False
    for sh in s.shapes:
        if sh.has_text_frame and "설계형 종신보험 검토 ·" in sh.text_frame.text:
            set_text(sh, f"설계형 종신보험 검토 · {i} / {N}"); done = True
        elif sh.name == "FOOTER_PAGE":
            tf = sh.text_frame; tf.margin_left = tf.margin_right = Emu(0)
            p = tf.paragraphs[0]; r = p.add_run(); r.text = f"설계형 종신보험 검토 · {i} / {N}"
            r.font.name = "JetBrains Mono"; r.font.size = Pt(9); r.font.color.rgb = rgb(GRAY); done = True
prs.save(DST)
print("saved", DST, "slides", N)
